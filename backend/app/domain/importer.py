"""Import existing customers from the previous system: vehicles, contacts, running monthly passes
and opening balances (dues or credit).

Two steps, same code path:
  plan(db, rows)          -> per-row validation and the actions that would be taken (dry run)
  apply(db, plan, user)   -> performs the actions of every row without errors

Re-importing the same file is safe: existing vehicles are updated rather than duplicated, a pass that
already covers the period is not created again, and an opening balance is posted only once per vehicle.

Imported passes were paid in the old system, so they create no ledger entries and do not count as pass
sales in reports (channel "IMPORT"). Opening balances are ledger ADJUSTMENTs with a fixed reason.
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any, Optional

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..db import utcnow
from ..models import LedgerEntry, LedgerKind, Pass, PassType, User, Vehicle, VehicleClass
from . import ledger, plates
from .lookup import site_tz
from .passes import _add_months
from .settings import get_setting

OPENING_REASON = "Opening balance (import) from previous system"

COLUMNS = ["plate", "vehicle_class", "name", "phone", "pass_type", "pass_start", "pass_end", "pass_amount",
           "opening_balance", "notes"]
HEADER_ALIASES = {
    "vehicle": "plate", "vehicle_number": "plate", "registration": "plate", "number_plate": "plate", "plate_number": "plate",
    "class": "vehicle_class", "type": "vehicle_class",
    "mobile": "phone", "mobile_number": "phone", "phone_number": "phone", "contact": "phone",
    "customer": "name", "customer_name": "name", "owner": "name",
    "pass": "pass_type", "pass_from": "pass_start", "valid_from": "pass_start", "start_date": "pass_start",
    "pass_to": "pass_end", "valid_till": "pass_end", "valid_to": "pass_end", "end_date": "pass_end",
    "amount": "pass_amount", "amount_paid": "pass_amount", "pass_price": "pass_amount",
    "balance": "opening_balance", "dues": "opening_balance", "due": "opening_balance",
    "remarks": "notes", "note": "notes",
}

TEMPLATE_ROWS = [
    {"plate": "MH43AB1234", "vehicle_class": "BIKE", "name": "Ravi Kumar", "phone": "9876543210", "pass_type": "Monthly",
     "pass_start": "2026-09-01", "pass_end": "2026-09-30", "pass_amount": "500", "opening_balance": "", "notes": ""},
    {"plate": "MH12CD5678", "vehicle_class": "BIKE", "name": "", "phone": "9123456780", "pass_type": "", "pass_start": "",
     "pass_end": "", "pass_amount": "", "opening_balance": "40", "notes": "Rs 40 unpaid in old register"},
]


class ImportFileError(ValueError):
    pass


# ------------------------------------------------------------------ reading files
def _norm_header(h: Any) -> str:
    k = re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(k, k)


def read_rows(content: bytes, filename: str) -> list[dict[str, Any]]:
    """Rows of a CSV or Excel (.xlsx) file as dicts keyed by the normalised column names."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001 - any parse failure is the user's file
            raise ImportFileError(f"could not read the Excel file: {e}") from e
        it = wb.active.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            raise ImportFileError("the sheet is empty")
        keys = [_norm_header(h) for h in header]
        rows = [dict(zip(keys, r)) for r in it]
    elif name.endswith((".csv", ".txt")) or not name:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.reader(io.StringIO(text))
        header = next(reader, None)
        if not header:
            raise ImportFileError("the file is empty")
        keys = [_norm_header(h) for h in header]
        rows = [dict(zip(keys, r)) for r in reader]
    else:
        raise ImportFileError("upload a .csv or .xlsx file")
    if "plate" not in keys:
        raise ImportFileError("the first row must contain column names, including 'plate'")
    # drop completely empty lines
    return [r for r in rows if any(str(v).strip() for v in r.values() if v is not None)]


def template_csv() -> str:
    """Column headers plus two example rows (a running monthly pass and a customer with old dues)."""
    today = utcnow().astimezone(site_tz()).date()
    start = today.replace(day=1)
    end = _add_months(datetime.combine(start, time(0)), 1).date() - timedelta(days=1)
    rows = [dict(r) for r in TEMPLATE_ROWS]
    rows[0].update(pass_start=start.isoformat(), pass_end=end.isoformat())
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS)
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


# ------------------------------------------------------------------ value parsing
def _s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v: Any) -> Optional[date]:
    if v is None or _s(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _s(v)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"date '{t}' not understood (use YYYY-MM-DD or DD-MM-YYYY)")


def _rupees_to_paise(v: Any) -> Optional[int]:
    t = _s(v).replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "").strip()
    if t == "":
        return None
    try:
        return int(round(float(t) * 100))
    except ValueError:
        raise ValueError(f"amount '{_s(v)}' is not a number") from None


# ------------------------------------------------------------------ planning
@dataclass
class RowPlan:
    row: int  # 1-based spreadsheet row number (header = row 1)
    plate: str = ""
    display_plate: str = ""
    status: str = "ok"  # ok | warning | error
    messages: list[str] = field(default_factory=list)
    actions: list[str] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict)

    def warn(self, m: str) -> None:
        self.messages.append(m)
        if self.status == "ok":
            self.status = "warning"

    def error(self, m: str) -> None:
        self.messages.append(m)
        self.status = "error"

    def as_dict(self) -> dict:
        return {"row": self.row, "plate": self.plate, "display_plate": self.display_plate, "status": self.status,
                "messages": self.messages, "actions": self.actions}


def plan(db: Session, rows: list[dict[str, Any]]) -> list[RowPlan]:
    tz = site_tz()
    today = utcnow().astimezone(tz).date()
    codes = get_setting(db, "state_codes")
    classes = {c.code: c for c in db.scalars(select(VehicleClass)).all()}
    pass_types = db.scalars(select(PassType)).all()
    seen: dict[str, int] = {}
    out: list[RowPlan] = []
    for i, r in enumerate(rows, start=2):
        rp = RowPlan(row=i)
        out.append(rp)
        raw_plate = _s(r.get("plate"))
        if not raw_plate:
            rp.error("plate is missing")
            continue
        corr = plates.correct(raw_plate, codes)
        if not corr.valid:
            rp.error(f"'{raw_plate}' is not a valid Indian number plate")
            continue
        rp.plate, rp.display_plate = corr.plate, plates.display(corr.plate)
        if corr.substitutions:
            rp.warn(f"plate read as {rp.display_plate} (from '{raw_plate}')")
        if rp.plate in seen:
            rp.error(f"same plate as row {seen[rp.plate]}")
            continue
        seen[rp.plate] = i

        vclass = (_s(r.get("vehicle_class")) or "BIKE").upper()
        vclass = {"TWO_WHEELER": "BIKE", "2W": "BIKE", "SCOOTER": "BIKE", "MOTORCYCLE": "BIKE", "4W": "CAR"}.get(vclass, vclass)
        if vclass not in classes:
            rp.error(f"vehicle class '{vclass}' unknown (use BIKE or CAR)")
            continue

        phone = None
        if _s(r.get("phone")):
            digits = re.sub(r"\D", "", _s(r.get("phone")))
            if len(digits) == 12 and digits.startswith("91"):
                digits = digits[2:]
            if len(digits) != 10:
                rp.error(f"mobile '{_s(r.get('phone'))}' must have 10 digits")
                continue
            phone = digits
        name = _s(r.get("name")) or None
        notes = _s(r.get("notes")) or None

        veh = db.scalars(select(Vehicle).where(Vehicle.plate == rp.plate)).first()
        if veh is None:
            rp.actions.append("new vehicle")
        else:
            if veh.vehicle_class != vclass:
                rp.error(f"already registered as {veh.vehicle_class}, file says {vclass}")
                continue
            if phone and veh.phone != phone:
                rp.actions.append("update mobile" if veh.phone else "add mobile")
                if veh.phone:
                    rp.warn(f"mobile changes from {veh.phone} to {phone}")
            if name and veh.name != name:
                rp.actions.append("update name")
        rp.data.update(vehicle_class=vclass, phone=phone, name=name, notes=notes)

        # ---- pass
        try:
            pass_type_name = _s(r.get("pass_type"))
            start = _date(r.get("pass_start"))
            end = _date(r.get("pass_end"))
            amount = _rupees_to_paise(r.get("pass_amount"))
            opening = _rupees_to_paise(r.get("opening_balance"))
        except ValueError as e:
            rp.error(str(e))
            continue
        if pass_type_name or end:
            pt = next((p for p in pass_types if p.vehicle_class == vclass and p.name.lower() == pass_type_name.lower()), None) \
                if pass_type_name else next((p for p in pass_types if p.vehicle_class == vclass and p.is_default), None)
            if pt is None:
                names = ", ".join(sorted({p.name for p in pass_types if p.vehicle_class == vclass}))
                rp.error(f"pass type '{pass_type_name}' not found for {vclass} (available: {names})")
                continue
            start = start or today
            if end is None:
                end_excl = _add_months(datetime.combine(start, time(0)), pt.period_value).date() \
                    if pt.period_unit == "MONTH" else start + timedelta(days=pt.period_value)
                end = end_excl - timedelta(days=1)
            if end < start:
                rp.error("pass end is before pass start")
                continue
            if end < today:
                rp.warn(f"pass ended on {end.isoformat()}: not imported")
            else:
                starts_at = datetime.combine(start, time(0), tzinfo=tz)
                ends_at = datetime.combine(end + timedelta(days=1), time(0), tzinfo=tz)
                overlap = veh is not None and db.scalar(select(func.count(Pass.id)).where(
                    Pass.vehicle_id == veh.id, Pass.status == "ACTIVE", Pass.starts_at < ends_at, Pass.ends_at > starts_at))
                if overlap:
                    rp.warn("already has a pass for this period: pass not imported again")
                else:
                    rp.actions.append(f"{pt.name} pass {start.isoformat()} to {end.isoformat()}")
                    rp.data.update(pass_type_id=pt.id, pass_starts_at=starts_at, pass_ends_at=ends_at,
                                   pass_amount=amount if amount is not None else pt.price_paise)

        # ---- opening balance
        if opening:
            already = veh is not None and db.scalar(select(func.count(LedgerEntry.id)).where(
                LedgerEntry.vehicle_id == veh.id, LedgerEntry.kind == LedgerKind.ADJUSTMENT,
                LedgerEntry.reason.like(OPENING_REASON + "%")))
            if already:
                rp.warn("opening balance was already imported: not added again")
            else:
                label = f"dues ₹{opening / 100:g}" if opening > 0 else f"credit ₹{-opening / 100:g}"
                rp.actions.append(f"opening {label}")
                rp.data["opening_paise"] = opening
        if not rp.actions and rp.status == "ok":
            rp.warn("nothing to change")
    return out


def summary(plans: list[RowPlan]) -> dict:
    ok = [p for p in plans if p.status != "error"]
    return {
        "rows": len(plans),
        "importable": len(ok),
        "errors": sum(1 for p in plans if p.status == "error"),
        "warnings": sum(1 for p in plans if p.status == "warning"),
        "new_vehicles": sum(1 for p in ok if "new vehicle" in p.actions),
        "contacts": sum(1 for p in ok if any(a.endswith("mobile") or a == "update name" for a in p.actions)),
        "passes": sum(1 for p in ok if "pass_type_id" in p.data),
        "opening_balances": sum(1 for p in ok if "opening_paise" in p.data),
        "opening_dues_paise": sum(p.data.get("opening_paise", 0) for p in ok if p.data.get("opening_paise", 0) > 0),
        "opening_credit_paise": -sum(p.data.get("opening_paise", 0) for p in ok if p.data.get("opening_paise", 0) < 0),
    }


# ------------------------------------------------------------------ applying
def apply(db: Session, plans: list[RowPlan], user: User) -> dict:
    batch = uuid.uuid4().hex[:8]
    now = utcnow()
    done = 0
    for rp in plans:
        if rp.status == "error" or not rp.actions:
            continue
        d = rp.data
        veh = db.scalars(select(Vehicle).where(Vehicle.plate == rp.plate)).first()
        if veh is None:
            veh = Vehicle(plate=rp.plate, plate_canon=plates.canonical(rp.plate), display_plate=rp.display_plate,
                          vehicle_class=d["vehicle_class"], first_seen=now, last_seen=now, balance_paise=0)
            db.add(veh)
            db.flush()
        if d.get("phone"):
            veh.phone = d["phone"]
        if d.get("name"):
            veh.name = d["name"]
        if d.get("notes"):
            veh.notes = (f"{veh.notes}\n" if veh.notes else "") + f"[import {batch}] {d['notes']}"
        if "pass_type_id" in d:
            db.add(Pass(vehicle_id=veh.id, vehicle_class=veh.vehicle_class, pass_type_id=d["pass_type_id"],
                        starts_at=d["pass_starts_at"], ends_at=d["pass_ends_at"], amount_paise=d["pass_amount"],
                        status="ACTIVE", sold_by=user.id, channel="IMPORT"))
        if d.get("opening_paise"):
            ledger.post(db, veh.id, LedgerKind.ADJUSTMENT, d["opening_paise"],
                        reason=f"{OPENING_REASON} [batch {batch}]", user_id=user.id)
        done += 1
    db.flush()
    return {"batch": batch, "imported_rows": done}

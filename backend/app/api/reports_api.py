"""Dashboard live view and reports (JSON, Excel, PDF)."""
from __future__ import annotations

import io
from datetime import timedelta
from typing import Any, Callable, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..adapters.gateway import get_gateway
from ..db import utcnow
from ..domain import cash, reports
from ..domain.lookup import local_date
from ..domain.settings import get_setting
from ..models import AnprEvent, Device, EventStatus, Setting, User
from .deps import get_db, supervisor
from .serial import event_json, iso

router = APIRouter(prefix="/api")


@router.get("/dashboard/live")
def live(db: Session = Depends(get_db), user: User = Depends(supervisor)):
    now = utcnow()
    today = local_date(now)
    latest = db.scalars(select(AnprEvent).where(AnprEvent.status != EventStatus.DUPLICATE)
                        .order_by(AnprEvent.ts.desc()).limit(30)).all()
    devs = []
    for d in db.scalars(select(Device).order_by(Device.kind, Device.id)).all():
        age = (now - d.last_seen).total_seconds() if d.last_seen else None
        devs.append({"id": d.id, "kind": d.kind, "gate_id": d.gate_id, "name": d.name, "last_seen": iso(d.last_seen),
                     "online": age is not None and age < 60, "age_s": age, "metrics": d.metrics})
    net = db.get(Setting, "_internet_status")
    return {
        "now": now.isoformat(), "date": today,
        "occupancy": reports.occupancy(db),
        "movements": reports.movements_per_hour(db, today),
        "latest_events": [event_json(e) for e in latest],
        "devices": devs,
        "internet": net.value if net else {"online": None},
        "gateway": {"name": get_gateway().name, "online": (net.value or {}).get("gateway_online") if net else None},
        "review_counts": reports.review_queue_counts(db),
        "unpaid_flagged": len(reports.unpaid_flagged_sessions(db)),
        "cash": {"holdings": cash.all_holdings(db), "reconciliation": cash.cash_reconciliation(db, today)},
        "revenue_today": (reports.daily_revenue(db, today) or [{}])[0],
    }


def _sheet_rows(data: Any) -> list[dict]:
    if isinstance(data, list):
        return data
    if isinstance(data, dict):
        for k in ("rows", "defaulters", "per_worker"):
            if isinstance(data.get(k), list):
                return data[k]
        return [{"key": k, "value": v} for k, v in data.items() if not isinstance(v, (list, dict))]
    return []


REPORTS: dict[str, Callable[..., Any]] = {
    "revenue": lambda db, s, e: reports.daily_revenue(db, s, e),
    "cash-reconciliation": lambda db, s, e: cash.cash_reconciliation(db, s),
    "worker-comparison": lambda db, s, e: reports.worker_comparison(db, s, e),
    "collections": lambda db, s, e: reports.collections_per_worker(db, s, e),
    "disputes": lambda db, s, e: reports.disputes_by_worker(db, s, e),
    "unpaid-sessions": lambda db, s, e: reports.sessions_without_payment(db, s, e),
    "passes": lambda db, s, e: reports.pass_report(db, s, e),
    "anpr-accuracy": lambda db, s, e: reports.anpr_accuracy(db, s, e),
    "peak-hours": lambda db, s, e: reports.peak_hours(db, s, e),
    "overrides": lambda db, s, e: reports.overrides_report(db, s, e),
    "defaulters": lambda db, s, e: reports.defaulters(db, int(get_setting(db, "alert_balance_threshold_paise"))),
}


@router.get("/reports/{name}")
def report(name: str, start: Optional[str] = None, end: Optional[str] = None, format: str = "json",
           db: Session = Depends(get_db), user: User = Depends(supervisor)):
    fn = REPORTS.get(name)
    if fn is None:
        raise HTTPException(404, f"unknown report; choose from {sorted(REPORTS)}")
    s = start or local_date(utcnow())
    data = fn(db, s, end)
    if format == "json":
        return data
    rows = _sheet_rows(data)
    title = f"{name} {s}{' to ' + end if end else ''}"
    if format == "xlsx":
        return _xlsx(rows, name, title)
    if format == "pdf":
        return _pdf(rows, name, title)
    raise HTTPException(400, "format must be json, xlsx or pdf")


def _flat(v: Any) -> Any:
    if isinstance(v, (list, dict)):
        return ", ".join(map(str, v)) if isinstance(v, list) else str(v)
    return v


def _xlsx(rows: list[dict], name: str, title: str) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = name[:30]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    cols = list(rows[0].keys()) if rows else []
    ws.append(cols)
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_flat(r.get(c)) for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(2, i).column_letter].width = max(10, min(40, len(c) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


def _pdf(rows: list[dict], name: str, title: str) -> StreamingResponse:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    cols = list(rows[0].keys())[:14] if rows else []
    data = [cols] + [[str(_flat(r.get(c)) if r.get(c) is not None else "")[:28] for c in cols] for r in rows]
    story = [Paragraph(title, getSampleStyleSheet()["Title"])]
    if data and cols:
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]))
        story.append(t)
    else:
        story.append(Paragraph("No data", getSampleStyleSheet()["Normal"]))
    doc.build(story)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'})
